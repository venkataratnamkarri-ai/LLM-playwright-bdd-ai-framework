"""
Excel Persistence & Cohere RAG question answering for Demoblaze laptop data.
"""

from __future__ import annotations
import json
import os
from pathlib import Path
import cohere
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook

HEADERS = ("Name", "Price")
ANSWER_KEYS = ("Name", "Price")

load_dotenv()


def write_laptops_excel(laptops: list[dict], output_path: str | Path) -> Path:
    if not laptops:
        raise ValueError("Cannot write an empty laptop catalog.")
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Laptops"
    sheet.append(HEADERS)
    for laptop in laptops:
        sheet.append((laptop["Name"], laptop["Price"]))
    sheet.freeze_panes = "A2"
    workbook.save(path)
    return path


def read_laptops_excel(path: str | Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    if headers != HEADERS:
        raise ValueError(f"Expected Excel columns {HEADERS}, got {headers}")
    laptops = []
    for name, price in sheet.iter_rows(min_row=2, values_only=True):
        if name is not None and price is not None:
            laptops.append({"Name": str(name), "Price": float(price)})
    workbook.close()
    if not laptops:
        raise ValueError("The Excel workbook contains no laptop rows.")
    return laptops


def laptops_context(laptops: list[dict]) -> str:
    return "\n".join(
        f"Name: {laptop['Name']} | Price: ${float(laptop['Price']):.2f}"
        for laptop in laptops
    )


class LaptopQuestionAnswerer:
    def __init__(self, api_key: str | None = None, model: str = "command-a-03-2025"):
        key = api_key or os.getenv("COHERE_API_KEY")
        if not key:
            raise ValueError("COHERE_API_KEY not found in .env")
        self.client = cohere.ClientV2(api_key=key)
        self.model = model

    def cheapest_laptop(self, laptops: list[dict]) -> dict:
        context = laptops_context(laptops)
        prompt = f"""Use only the laptop catalog below. Select the laptop with the lowest numeric price.
If prices tie, select the first listed. Return exactly one JSON object with keys:
Name (string) and Price (number).

Catalog:
{context}

Return valid JSON only.
"""
        response = self.client.chat(
            model=self.model,
            messages=[{"role": "user", "content": prompt}],
            temperature=0,
            max_tokens=150,
        )
        text = "".join(
            item.text for item in response.message.content
            if getattr(item, "text", None)
        )
        return validate_cheapest_laptop(text, laptops)


def validate_cheapest_laptop(response: str, laptops: list[dict]) -> dict:
    candidate = response.strip()
    fence = chr(96) * 3
    candidate = candidate.removeprefix(fence + "json").removeprefix(fence).removesuffix(fence).strip()
    try:
        answer = json.loads(candidate)
    except json.JSONDecodeError as exc:
        raise AssertionError(f"Cohere response was not valid JSON: {response!r}") from exc

    expected = min(laptops, key=lambda laptop: float(laptop["Price"]))
    if answer["Name"].strip() != expected["Name"]:
        raise AssertionError(f"Expected cheapest laptop {expected['Name']!r}; got {answer['Name']!r}")
    return answer
